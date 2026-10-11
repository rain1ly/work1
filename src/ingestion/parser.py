"""M2 解析与 OCR 模块。

负责把不同格式的文件转成纯文本。
- .md / .txt  直接读取（front-matter 由 loader 处理）
- .pdf        优先 PyMuPDF；未安装则给出明确警告，不静默失败
- .docx       优先 python-docx
- .xlsx       优先 openpyxl，逐行拍平成文本
- 扫描件     预留 OCR 接口（PaddleOCR），未安装时在 warnings 中提示

设计原则：解析失败必须显式报出来，不能悄悄跳过 —— 否则知识库里缺了文件都不知道。
"""

from __future__ import annotations

from pathlib import Path


def _read_text(path: Path) -> str:
    for enc in ("utf-8", "utf-8-sig", "gbk", "gb18030"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def _parse_pdf(path: Path, warnings: list[str]) -> str:
    try:
        import fitz  # PyMuPDF
    except Exception:
        warnings.append(f"[{path.name}] 未安装 PyMuPDF，无法解析 PDF（pip install pymupdf）")
        return ""
    pages: list[str] = []
    with fitz.open(path) as doc:
        for i, page in enumerate(doc, 1):
            text = page.get_text("text").strip()
            if not text:
                warnings.append(
                    f"[{path.name}] 第 {i} 页无可提取文本，疑似扫描件，需接入 OCR（PaddleOCR）"
                )
                continue
            pages.append(text)
    return "\n\n".join(pages)


def _parse_docx(path: Path, warnings: list[str]) -> str:
    try:
        import docx  # python-docx
    except Exception:
        warnings.append(f"[{path.name}] 未安装 python-docx，无法解析 Word（pip install python-docx）")
        return ""
    d = docx.Document(str(path))
    blocks = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


def _parse_xlsx(path: Path, warnings: list[str]) -> str:
    try:
        import openpyxl  # type: ignore
    except Exception:
        warnings.append(f"[{path.name}] 未安装 openpyxl，无法解析 Excel（pip install openpyxl）")
        return ""
    wb = openpyxl.load_workbook(str(path), data_only=True)
    blocks: list[str] = []
    for ws in wb.worksheets:
        blocks.append(f"## 工作表：{ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


def ocr_image(path: Path, warnings: list[str]) -> str:
    """扫描件 OCR 接口（P2）。未安装 PaddleOCR 时明确提示。"""
    try:
        from paddleocr import PaddleOCR  # type: ignore
    except Exception:
        warnings.append(f"[{path.name}] 未安装 PaddleOCR，跳过 OCR（pip install paddleocr paddlepaddle）")
        return ""
    ocr = PaddleOCR(use_angle_cls=True, lang="ch", show_log=False)
    result = ocr.ocr(str(path), cls=True)
    lines: list[str] = []
    for page in result or []:
        for item in page or []:
            try:
                lines.append(item[1][0])
            except Exception:
                continue
    return "\n".join(lines)


def parse(path: Path) -> tuple[str, list[str]]:
    """返回 (纯文本, 警告列表)。"""
    warnings: list[str] = []
    ext = path.suffix.lower()
    if ext in (".md", ".txt"):
        text = _read_text(path)
    elif ext == ".pdf":
        text = _parse_pdf(path, warnings)
    elif ext == ".docx":
        text = _parse_docx(path, warnings)
    elif ext == ".xlsx":
        text = _parse_xlsx(path, warnings)
    elif ext in (".png", ".jpg", ".jpeg"):
        text = ocr_image(path, warnings)
    else:
        warnings.append(f"[{path.name}] 不支持的格式：{ext}")
        text = ""
    return text, warnings
